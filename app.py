import asyncio
import logging
import sys
import sqlite3
import pandas as pd
from openpyxl.styles import Font
from indexer import GooglePhotosIndexer
from database import Database

logging.basicConfig(
    level=logging.INFO,
    format="[%(asctime)s] %(levelname)s - %(message)s",
    datefmt="%H:%M:%S"
)

ALBUM_FILE = "album_links.txt"

def load_all_albums(file_path: str) -> dict[str, str]:
    albums = {}
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                if ":" in line:
                    name, url = line.split(":", 1)
                    albums[name.strip()] = url.strip().strip('"').strip("'")
    except FileNotFoundError:
        print(f"Error: {file_path} not found.")
    return albums

def export_to_excel() -> None:
    db_name = "google_photos.db"
    excel_name = "google_photos_export.xlsx"
    
    print(f"\n--- Generating Excel Export ---")
    try:
        conn = sqlite3.connect(db_name)
        
        with pd.ExcelWriter(excel_name, engine='openpyxl') as writer:
            # 1. Output the raw DataFrames to sheets
            size_df = pd.read_sql_query("SELECT * FROM vw_album_storage", conn)
            size_df.to_excel(writer, sheet_name='Storage Size', index=False)
            
            albums_df = pd.read_sql_query("SELECT * FROM albums", conn)
            albums_df.to_excel(writer, sheet_name='Albums', index=False)
            
            media_df = pd.read_sql_query("SELECT * FROM vw_media_formatted", conn)
            media_df.to_excel(writer, sheet_name='Media', index=False)
            
            workbook = writer.book
            
            # 2. Fix Column Widths
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                for col in worksheet.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            val_str = str(cell.value)
                            if len(val_str) > max_length:
                                max_length = len(val_str)
                        except:
                            pass
                    worksheet.column_dimensions[column].width = min((max_length + 2), 50)
            
            # 3. Apply Native Hyperlinks
            media_sheet = workbook['Media']
            link_col_idx = None
            
            for idx, cell in enumerate(media_sheet[1], 1):
                if cell.value == 'Share Link':
                    link_col_idx = idx
                    break
            
            if link_col_idx:
                link_font = Font(color="0563C1", underline="single")
                for row in range(2, media_sheet.max_row + 1):
                    cell = media_sheet.cell(row=row, column=link_col_idx)
                    url = cell.value
                    
                    if url and isinstance(url, str) and url.startswith("http"):
                        cell.hyperlink = url
                        cell.value = "Open in Photos"
                        cell.font = link_font
                    else:
                        cell.value = "" 
            
        print(f"Success! Export saved to {excel_name}")
        
    except PermissionError:
        print(f"\nERROR: The file '{excel_name}' is open. Close it and run export again.")
    except Exception as e:
        print(f"\nExcel export failed: {e}")
    finally:
        if 'conn' in locals():
            conn.close()

async def process_album(album_name: str, url: str, db: Database) -> None:
    print(f"\n--- Indexing Album: {album_name} ---")
    print(f"URL: {url}")
    
    indexer = GooglePhotosIndexer()
    try:
        result = await indexer.index(url)
        print("\nSuccess!")
        print(f"Album ID: {result.album.id}")
        print(f"Title:    {result.album.title}")
        print(f"Media:    {result.media_count} items stored")
        print(f"Time:     {result.elapsed_seconds:.2f}s")
        
    except Exception as e:
        print(f"\nError processing {album_name}: {e}")

async def main() -> None:
    albums = load_all_albums(ALBUM_FILE)
    if not albums:
        return
        
    db = Database()
    targets = {}

    if len(sys.argv) > 1:
        target_name = sys.argv[1]
        if target_name in albums:
            targets[target_name] = albums[target_name]
        else:
            print(f"Error: Album '{target_name}' not found in {ALBUM_FILE}.")
            return
    else:
        targets = albums
        print(f"No specific album requested. Processing all {len(targets)} albums...")

    for name, url in targets.items():
        await process_album(name, url, db)
        
    export_to_excel()

if __name__ == "__main__":
    asyncio.run(main())